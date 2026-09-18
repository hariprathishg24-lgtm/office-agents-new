from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents" / "agents-office-unattended-operation-runbook.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for n, size in [('Title', 25), ('Subtitle', 12), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 10.5)]:
    s = styles[n]; s.font.name = 'Aptos Display' if n != 'Normal' else 'Aptos'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
    s.font.bold = n != 'Subtitle'; s.paragraph_format.space_before = Pt(12 if n != 'Title' else 0); s.paragraph_format.space_after = Pt(5)
styles['Title'].paragraph_format.space_after = Pt(10)
# Word's built-in Title style may carry a theme-colored bottom border. This runbook uses spacing only.
title_ppr = styles['Title']._element.get_or_add_pPr()
for old in list(title_ppr.findall(qn('w:pBdr'))): title_ppr.remove(old)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    el = tblPr.find(qn('w:tblBorders'))
    if el is None: el = OxmlElement('w:tblBorders'); tblPr.append(el)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement('w:'+edge); tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'6'); tag.set(qn('w:color'),color); el.append(tag)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)

def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr(); flag = OxmlElement('w:cantSplit'); flag.set(qn('w:val'),'true'); trPr.append(flag)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'203864'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True; r.font.size=Pt(8.5)
    repeat_header(t.rows[0])
    keep_row_together(t.rows[0])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        keep_row_together(t.rows[-1])
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F3F6FA')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(8.3)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    for row in t.rows:
        for c in row.cells: set_cell_margin(c)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def bullets(items, numbered=False):
    for i,x in enumerate(items, 1):
        if numbered:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25); p.paragraph_format.space_after=Pt(3); p.add_run(f'{i}.  '); p.add_run(x)
        else:
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(x)

def checklist(items):
    for x in items:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.12); p.paragraph_format.first_line_indent=Inches(-.12); p.add_run('☐  ').bold=True; p.add_run(x)

page_break_count = 0
def page_break():
    global page_break_count
    page_break_count += 1
    if page_break_count == 1: doc.add_page_break()

# Cover
p=doc.add_paragraph(style='Title'); p.add_run('Agents Office Unattended Operation Runbook')
p=doc.add_paragraph(style='Subtitle'); p.add_run('A controlled path from the current verified build to dependable daily operation')
doc.add_paragraph('Origin Pixel Solutions  |  Version 1.0  |  18 September 2026')
doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run('Current conclusion'); r.bold=True
doc.add_paragraph('The software is ready for a supervised pilot. It is not ready for unattended production until the owner completes the operating decisions, reviews the active roles, observes a limited real pilot, and verifies restart and monitoring on this Windows machine.')
table(['Gate','Current result','Meaning'],[
    ('A  Deterministic reliability','Complete','97 of 97 automated tests across 26 suites.'),
    ('B  Integrated local workflow','Complete','38 of 38 checks and 7 of 7 browser drills.'),
    ('C  Limited real operation','Not started','Requires owner-approved limits, connectors, role reviews and an observed pilot.'),
], [1.35,1.55,4.05])
doc.add_paragraph('This runbook defines safe unattended operation. Agents may perform approved read-only and local work on schedule. Sending, publishing, paying, signing, deleting, or changing an external system continues to require approval for the exact recipient and payload.')
doc.add_heading('How to use this runbook', level=1)
bullets([
    'Complete Stages 1 through 6 in order. Do not install autostart before the pilot passes.',
    'Record every owner decision in the tables provided. Do not substitute assumptions for missing limits.',
    'Stop at any failed gate, preserve evidence, fix the cause, and repeat that gate before proceeding.',
    'Keep acquisition and research inactive until their limits are recorded and their connectors are ready.'
])

page_break()
doc.add_heading('1  Current verified state', level=1)
doc.add_paragraph('This baseline prevents work from being declared complete merely because the office contains 119 seats or because automated tests pass.')
table(['Area','Verified state','Remaining requirement'],[
    ('Code and tests','Clean reliability-handoff branch; Gates A and B pass.','Preserve the clean baseline and tag the pilot candidate.'),
    ('Roles','119 seats; 53 generic, 52 briefed, 14 contracted, 0 tested.','Owner reviews the 14 contracted seats before they run unattended.'),
    ('Task store','29 tasks: 18 done, 7 cancelled, 3 waiting, 1 interrupted.','Resolve two stale tasks, inspect the Instagram draft and reconcile the interrupted CEO run.'),
    ('Routines','Five enabled and two paused.','Confirm every enabled routine has a tested owner role and acceptable schedule.'),
    ('Acquisition','Inactive; four limits unset.','Set limits and connect prospect research before activation.'),
    ('Research','Inactive; two limits unset.','Set cadence and findings limit before activation.'),
    ('Service','Office down; autostart not registered.','Complete pilot, then verify scheduled startup and crash recovery.'),
    ('Connectors','Apollo needs authorization; Zoho reports an OAuth scope mismatch.','Authorize only required connectors and verify least privilege.'),
], [1.25,2.45,3.25])

doc.add_heading('Definition of unattended operation', level=2)
doc.add_paragraph('Unattended means the local service can start, stay healthy, run bounded read-only or local routines, produce drafts, pause on limits, and recover safely without someone keeping the page open. It does not mean autonomous authority over external actions.')
table(['May run without approval','Must wait for owner approval'],[
    ('Read current notes; research public sources; create local prospect records; produce drafts; run deterministic checks; prepare reports.','Send email or messages; publish content; purchase tools or ads; sign or accept terms; delete or modify external records; approve another role’s work.'),
], [3.45,3.5])

page_break()
doc.add_heading('2  Stage 1  Record operating decisions', level=1)
doc.add_paragraph('These values turn open-ended work into bounded work. Recommended pilot values are intentionally conservative.')
table(['Decision','Recommended pilot value','Owner value','Reason'],[
    ('Outreach drafts per day','5','____________','Small enough for daily factual review.'),
    ('Follow-up delay','4 days','____________','Avoids immediate repeated contact.'),
    ('Maximum follow-ups','2','____________','Creates a hard stop.'),
    ('Spend ceiling','INR 0','____________','No paid tools or ads during the pilot.'),
    ('Research runs','1 per week','____________','Limits token and source-review load.'),
    ('Findings per run','5','____________','Keeps the digest reviewable.'),
    ('Pilot duration (manual)','7 calendar days','____________','No config field: the owner enforces this by ending the pilot.'),
    ('Pilot prospect cap (manual)','25 total','____________','No config field: only the four acquisition limits are enforced in code.'),
], [1.55,1.55,1.3,2.05])
doc.add_paragraph('Optional deadline: ____________________. Leave it unset rather than inventing one.')
doc.add_heading('Configuration procedure', level=2)
bullets([
    'Write the four acquisition values to brain/Agents Office/acquisition.json. Keep active false until Stage 5.',
    'Write runsPerWeek and findingsPerRun to brain/Agents Office/research.json. Keep active false until Stage 5.',
    'Commit the brain repository with a message that names the approved limits.',
    'Run npm run coverage and the isolated checks after editing.'
], numbered=True)
checklist([
    'All six required values are explicit and non-null.',
    'Spend ceiling uses the same currency understood by the owner.',
    'No connector or schedule was activated merely by recording limits.',
    'The owner signed the decision record at the end of this runbook.'
])

doc.add_heading('3  Stage 2  Prepare connectors and local security', level=1)
table(['Connector or boundary','Required action','Acceptance evidence'],[
    ('Apollo.io','Authorize the existing Claude connector with the minimum scopes needed for prospect research.','A read-only test returns one known company and its source; no message is sent.'),
    ('Zoho','Fix OAUTH_SCOPE_MISMATCH only if Zoho is required for the pilot; otherwise leave disabled.','A read-only identity or list operation succeeds with the intended account.'),
    ('Notion','Keep denied unless the owner separately changes policy.','Configuration continues to show Notion denied.'),
    ('Server binding','Keep 127.0.0.1 loopback binding.','Health reports this machine only; foreign Host requests are refused.'),
    ('Claude login','Confirm the intended user session and current plan headroom.','One observed read-only task completes without login or usage failure.'),
], [1.25,3.05,2.65])
doc.add_paragraph('Never paste credentials into company notes, task text, screenshots, logs, contracts or fixture results. Use the connector’s normal authorization flow.')

page_break()
doc.add_heading('4  Stage 3  Clear the queue safely', level=1)
doc.add_paragraph('Open work affects routine scheduling. Resolve it before the pilot so old drafts cannot be mistaken for fresh work.')
doc.add_paragraph('This is not housekeeping. A routine has at most one open run, and "waiting" counts as open, so outbound-first-touch and inbound-qualify have not fired since 17 September and will not fire again until their two drafts are closed. Two of the five enabled routines are dead until Stage 3 is done. Both drafts are "nothing found" reports written about twenty minutes before the office began filing such drafts as finished reports, and neither carries a TO: recipient: cancelling costs no Claude run, while approving spends one carrying out an outbound step that has no recipient and nothing to send.')
table(['Record','Required disposition','Evidence'],[
    ('mu4ybp1lzrs0','Cancel the stale first-touch draft after confirming nothing was sent.','Task state cancelled; history records the owner action.'),
    ('mu4yiaq6vr5p','Cancel the stale inbound qualification draft after confirming nothing was sent.','Task state cancelled; history records the owner action.'),
    ('mu4ybsp45ila','Read the four Instagram drafts. Approve, reject with feedback, or cancel.','Final state and review note are explicit.'),
    ('Interrupted CEO run','Restart under observation and confirm boot reconciliation creates no external action.','Attempt history shows interrupted then safely requeued or resolved.'),
], [1.45,3.4,2.1])
checklist([
    'No task remains doing from a previous process.',
    'Every waiting item has a named owner decision.',
    'Unknown external outcomes are reconciled before retry.',
    'Cancelling work does not delete its history.'
])

doc.add_heading('5  Stage 4  Review and activate role capability', level=1)
doc.add_paragraph('A passing fixture run is evidence for review, not permission to mark a role tested automatically. The owner must read the saved answers and approve the current contract hash.')
table(['Role group','Seats','Review focus'],[
    ('First-client path','pros, enzo, ilm, lexi, folo, cmail, piper, qa, dlead, pco','Sources, fit, truthful claims, correct prices, approval boundaries and handoff quality.'),
    ('Enabled routine roles','iggy, ceo','No fabricated proof; useful content; evidence-based executive counts; no executive approval authority.'),
    ('Paused routine roles','vertical, prodz','vertical: ICP changes remain proposals based on real replies and require owner review. prodz: offers are packaged, never priced — every figure comes from the offer ladder verbatim.'),
], [1.45,2.25,3.25])
doc.add_paragraph('Both paused routine roles were contracted on 18 September and pass 3/3, so neither can run on its job title if its routine is later unpaused. Adding a capability contract promotes a seat to "contracted", which is the bar readiness.requireForOutbound uses to permit a send without an owner override: never add a contract without running that seat’s three fixtures in the same pass, or send authority widens with no evidence behind it.')
doc.add_heading('Owner review method', level=2)
bullets([
    'Open the latest results file for one seat and confirm all three cases passed: normal, missing input and misleading input.',
    'Read the actual answer, not only the automated verdict. Confirm facts, prices, sources, tone, boundary handling and final state.',
    'Compare the behavior with that seat’s current capability contract.',
    'Approve through the coverage review endpoint or interface only when the evidence is acceptable. A contract edit invalidates an older review.',
    'Record corrections as feedback, rerun the seat and review the new evidence.'
], numbered=True)
checklist(['All 14 contracted seats show tested.','No review was recorded by an agent on the owner’s behalf.','Any rejected seat is excluded from enabled routines until it passes.'])

page_break()
doc.add_heading('6  Stage 5  Run Gate C as a supervised pilot', level=1)
doc.add_paragraph('The pilot proves real operation with narrow exposure. Keep autostart off. Start the office manually, keep the owner available, and inspect outcomes daily.')
table(['Pilot day','Action','Pass condition'],[
    ('Day 0','Back up task state, routines and brain; record commit IDs; start manually.','Health is ready, one instance owns the data lock, and no unexpected catch-up action occurs.'),
    ('Day 1','Run one read-only prospect research task through Apollo.','Sources are real and dated; zero results produces an honest no-op.'),
    ('Day 2','Allow up to five prospect records and drafts.','No duplicates, invented facts or sends; every draft waits for approval.'),
    ('Day 3','Approve one carefully reviewed message to a test or deliberately selected real recipient.','Exactly one send, correct recipient and payload, remote reference recorded.'),
    ('Day 4','Exercise follow-up timing without sending an unreviewed message.','Cadence and maximum follow-ups are respected.'),
    ('Day 5','Pause and resume the office; simulate network loss on a read-only task.','No work starts while paused; status explains the network failure; retry is bounded.'),
    ('Day 6','Perform a controlled restart during read-only work.','Task is recovered once; lock and attempt history remain correct.'),
    ('Day 7','Review metrics and decide pass, extend or roll back.','All mandatory thresholds below are met.'),
], [.8,3.25,3.0])
doc.add_heading('Pilot metrics and mandatory thresholds', level=2)
table(['Metric','Required threshold'],[
    ('Duplicate external actions','0'),('Unapproved external actions','0'),('Invented client, result, price or prospect facts','0'),
    ('Unknown outcomes left unreconciled','0'),('Task recovery duplication','0'),('Daily owner intervention','Recorded and acceptable'),
    ('Connector and Claude cost','Within approved ceiling'),('Critical test or health failure','0 unresolved'),
], [3.7,3.35])

page_break()
doc.add_heading('7  Stage 6  Enable unattended Windows operation', level=1)
doc.add_paragraph('Install autostart only after Gate C passes and the owner accepts the pilot report.')
doc.add_heading('Preinstallation checks', level=2)
checklist([
    'The office is stopped cleanly and no node process owns port 4520.',
    'data/office.lock is absent after clean shutdown.',
    'Five enabled routines have tested role owners or are paused.',
    'Acquisition and research active flags match the owner’s decision.',
    'The Claude login and required connectors belong to the intended Windows user.',
    'The current source and brain commits are recorded for rollback.'
])
doc.add_heading('Install and verify', level=2)
doc.add_paragraph('Run from the nested repository in an ordinary PowerShell session under the intended Windows account:')
for cmd in [
    r'powershell -ExecutionPolicy Bypass -File scripts\install-autostart.ps1',
    r'powershell -ExecutionPolicy Bypass -File scripts\install-autostart.ps1 -Check',
]:
    p=doc.add_paragraph(); p.style='Intense Quote'; p.add_run(cmd)
bullets([
    'Sign out and back in during a controlled window. Confirm exactly one office process starts.',
    'Open http://localhost:4520/api/health and /ops. Confirm readiness, heartbeat, routines and limits.',
    'Terminate the office process once to verify bounded restart backoff. Confirm a second copy cannot claim the same data folder.',
    'Use the Stop control to verify an intentional stop is not restarted.',
    'Confirm logs are written and remain readable after restart.'
], numbered=True)

doc.add_heading('Rollback', level=2)
doc.add_paragraph('If unattended behavior is unsafe, pause the office first, preserve task and attempt state, reconcile possible external actions, then remove autostart:')
p=doc.add_paragraph(); p.style='Intense Quote'; p.add_run(r'powershell -ExecutionPolicy Bypass -File scripts\install-autostart.ps1 -Uninstall')
doc.add_paragraph('Restore code only after preserving current state. Reverting a local database or Git commit does not prove that an external message was never sent.')

page_break()
doc.add_heading('8  Daily and weekly operating procedure', level=1)
table(['Frequency','Owner check','Action on failure'],[
    ('Daily','Heartbeat; failed tasks; waiting approvals; unknown outcomes; usage; next routines.','Pause if authority, duplication or corrupted state is uncertain.'),
    ('Daily','Review every outbound draft and exact recipient before approval.','Reject with specific feedback; never repair and send outside the approval flow.'),
    ('Weekly','Prospects, replies, source quality, costs, corrections and first-client movement.','Tighten limits or pause channels that create low-quality work.'),
    ('Weekly','Research findings and proposed playbook changes.','Reject unsupported, stale or policy-changing claims.'),
    ('Monthly','Contract hashes, role reviews, connector scopes, logs, backups and recovery drill.','Invalidate stale reviews; rotate/revoke unnecessary access.'),
], [1.05,3.35,2.55])
doc.add_heading('Safe response playbooks', level=2)
table(['Condition','Required response'],[
    ('Claude login expired','Pause dispatch; log in interactively; run one read-only check; resume only after success.'),
    ('Usage or rate limit','Let bounded read-only retries stop; do not retry outbound work; resume after the window resets.'),
    ('Network unavailable','Keep drafts local; surface the problem; retry only demonstrably read-only work with backoff.'),
    ('Outcome unknown','Do not resend. Check the connector’s sent folder or remote record, then reconcile in the office.'),
    ('Corrupt task store','Stop; preserve the corrupt file; use validated backup recovery; never replace it with an empty list.'),
    ('Lock or recovery marker','Inspect the owning process. Remove nothing merely because HTTP health is unavailable.'),
    ('Wrong or fabricated output','Reject, record correction, update contract or skill if recurring, rerun fixtures and review again.'),
], [1.6,5.35])

page_break()
doc.add_heading('9  Final production acceptance checklist', level=1)
doc.add_paragraph('Unattended operation is approved only when every mandatory item below is checked.')
checklist([
    'Gate A passes: all deterministic tests are green on the production candidate.',
    'Gate B passes: build, smoke checks and browser drill are green on the same candidate.',
    'All routine-owning roles that will run are tested against their current contracts.',
    'Apollo or the selected research source is authorized and verified read-only.',
    'Acquisition and research limits are explicit, reviewed and enforced.',
    'The task queue contains no unexplained doing state or unknown external outcome.',
    'Gate C completed with zero duplicate, unapproved or fabricated external actions.',
    'Pause, graceful stop, crash recovery, network failure and expired-login behavior were observed.',
    'Autostart points to the intended checkout and Windows user and starts exactly one instance.',
    'Health, heartbeat, logs, usage guard and routine status are visible in /ops.',
    'Rollback commits and state backups are recorded and recoverable.',
    'The owner accepts the residual limits of a laptop: sleep, shutdown, network and provider availability interrupt work.'
])
doc.add_heading('Production decision', level=2)
table(['Decision','Selection'],[
    ('Gate C result','☐ Pass   ☐ Extend pilot   ☐ Roll back'),
    ('Acquisition','☐ Active   ☐ Inactive'),
    ('Research','☐ Active   ☐ Inactive'),
    ('Autostart','☐ Install   ☐ Do not install'),
    ('Outbound authority','Exact-action owner approval remains required'),
], [2.4,4.55])
doc.add_paragraph('Owner name: ____________________________________')
doc.add_paragraph('Date: ____________________    Signature or recorded approval: ____________________________________')

doc.add_heading('10  Command reference', level=1)
doc.add_paragraph('Run all commands from C:\\Users\\harip\\Downloads\\Gittest\\agents-office-main\\agents-office-main. Keep real-operation commands separate from isolated checks. Do not set CHECK_LIVE unless a deliberate live test is authorized.')
table(['Purpose','Command'],[
    ('Automated tests',r'node --test --test-concurrency=1 test/*.test.mjs'),
    ('Integrated checks',r'$env:AO_BUILD_LOCAL_FILES="1"; $env:CHECK_LIVE="0"; node check.mjs'),
    ('Browser drill',r'node test/browser-drill.mjs'),
    ('Coverage',r'npm run coverage'),
    ('Start manually',r'start-office.cmd'),
    ('Check autostart',r'powershell -ExecutionPolicy Bypass -File scripts\install-autostart.ps1 -Check'),
    ('Remove autostart',r'powershell -ExecutionPolicy Bypass -File scripts\install-autostart.ps1 -Uninstall'),
], [2.0,4.95])

# Footer
for section in doc.sections:
    footer=section.footer
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Origin Pixel Solutions  |  Agents Office unattended operation runbook'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

doc.core_properties.title='Agents Office Unattended Operation Runbook'
doc.core_properties.subject='Operational readiness, pilot, deployment, monitoring and rollback'
doc.core_properties.author='Origin Pixel Solutions'
doc.save(OUT)
print(OUT)
